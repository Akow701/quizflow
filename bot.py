import os
import re
import sqlite3
import logging
import tempfile
from pathlib import Path

from telegram import Update, Poll, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application, CommandHandler, MessageHandler,
    PollAnswerHandler, CallbackQueryHandler,
    ContextTypes, filters
)

import pdfplumber
import docx
import openpyxl

logging.basicConfig(format="%(asctime)s - %(levelname)s - %(message)s", level=logging.INFO)
logger = logging.getLogger(__name__)

BOT_TOKEN = os.environ["BOT_TOKEN"]

# ─── DATABASE ────────────────────────────────────────────────────────────────

def init_db():
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS quizzes (
            quiz_id    TEXT PRIMARY KEY,
            owner_id   INTEGER,
            title      TEXT,
            questions  TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS sessions (
            session_id  TEXT PRIMARY KEY,
            chat_id     INTEGER,
            quiz_id     TEXT,
            q_index     INTEGER DEFAULT 0,
            active      INTEGER DEFAULT 1
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS scores (
            chat_id    INTEGER,
            user_id    INTEGER,
            username   TEXT,
            first_name TEXT,
            correct    INTEGER DEFAULT 0,
            total      INTEGER DEFAULT 0,
            PRIMARY KEY (chat_id, user_id)
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS active_polls (
            poll_id     TEXT PRIMARY KEY,
            chat_id     INTEGER,
            correct_idx INTEGER,
            session_id  TEXT
        )
    """)
    conn.commit()
    conn.close()

def save_quiz(quiz_id, owner_id, title, questions_json):
    import json
    conn = sqlite3.connect("quiz.db")
    conn.execute("INSERT OR REPLACE INTO quizzes VALUES (?,?,?,?,CURRENT_TIMESTAMP)",
                 (quiz_id, owner_id, title, json.dumps(questions_json, ensure_ascii=False)))
    conn.commit()
    conn.close()

def get_quiz(quiz_id):
    import json
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("SELECT quiz_id, owner_id, title, questions FROM quizzes WHERE quiz_id=?", (quiz_id,))
    row = c.fetchone()
    conn.close()
    if row:
        return {"quiz_id": row[0], "owner_id": row[1], "title": row[2], "questions": json.loads(row[3])}
    return None

def get_user_quizzes(owner_id):
    import json
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("SELECT quiz_id, title, questions FROM quizzes WHERE owner_id=? ORDER BY created_at DESC LIMIT 20", (owner_id,))
    rows = c.fetchall()
    conn.close()
    return [{"quiz_id": r[0], "title": r[1], "count": len(json.loads(r[2]))} for r in rows]

def delete_quiz(quiz_id):
    conn = sqlite3.connect("quiz.db")
    conn.execute("DELETE FROM quizzes WHERE quiz_id=?", (quiz_id,))
    conn.commit()
    conn.close()

def create_session(session_id, chat_id, quiz_id):
    conn = sqlite3.connect("quiz.db")
    conn.execute("INSERT OR REPLACE INTO sessions VALUES (?,?,?,0,1)", (session_id, chat_id, quiz_id))
    conn.commit()
    conn.close()

def get_session(session_id):
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("SELECT session_id, chat_id, quiz_id, q_index, active FROM sessions WHERE session_id=?", (session_id,))
    row = c.fetchone()
    conn.close()
    return row

def get_active_session(chat_id):
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("SELECT session_id, quiz_id, q_index FROM sessions WHERE chat_id=? AND active=1 ORDER BY rowid DESC LIMIT 1", (chat_id,))
    row = c.fetchone()
    conn.close()
    return row

def update_session_index(session_id, q_index):
    conn = sqlite3.connect("quiz.db")
    conn.execute("UPDATE sessions SET q_index=? WHERE session_id=?", (q_index, session_id))
    conn.commit()
    conn.close()

def close_session(session_id):
    conn = sqlite3.connect("quiz.db")
    conn.execute("UPDATE sessions SET active=0 WHERE session_id=?", (session_id,))
    conn.commit()
    conn.close()

def close_all_sessions(chat_id):
    conn = sqlite3.connect("quiz.db")
    conn.execute("UPDATE sessions SET active=0 WHERE chat_id=?", (chat_id,))
    conn.commit()
    conn.close()

def save_poll(poll_id, chat_id, correct_idx, session_id):
    conn = sqlite3.connect("quiz.db")
    conn.execute("INSERT OR REPLACE INTO active_polls VALUES (?,?,?,?)",
                 (poll_id, chat_id, correct_idx, session_id))
    conn.commit()
    conn.close()

def get_poll(poll_id):
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("SELECT chat_id, correct_idx, session_id FROM active_polls WHERE poll_id=?", (poll_id,))
    row = c.fetchone()
    conn.close()
    return row

def upsert_score(chat_id, user_id, username, first_name, is_correct):
    conn = sqlite3.connect("quiz.db")
    conn.execute("""
        INSERT INTO scores (chat_id, user_id, username, first_name, correct, total)
        VALUES (?, ?, ?, ?, ?, 1)
        ON CONFLICT(chat_id, user_id) DO UPDATE SET
            username=excluded.username, first_name=excluded.first_name,
            correct=correct+excluded.correct, total=total+1
    """, (chat_id, user_id, username, first_name, 1 if is_correct else 0))
    conn.commit()
    conn.close()

def get_leaderboard(chat_id, limit=10):
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("""
        SELECT first_name, username, correct, total FROM scores
        WHERE chat_id=? ORDER BY correct DESC, total ASC LIMIT ?
    """, (chat_id, limit))
    rows = c.fetchall()
    conn.close()
    return rows

def get_quiz_stats(quiz_id):
    conn = sqlite3.connect("quiz.db")
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM sessions WHERE quiz_id=?", (quiz_id,))
    sessions = c.fetchone()[0]
    conn.close()
    return sessions

# ─── FILE READERS ─────────────────────────────────────────────────────────────

def read_txt(path):
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()

def read_pdf(path):
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    return text

def read_docx(path):
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs)

def read_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = "\t".join(str(c) if c is not None else "" for c in row)
            if line.strip():
                lines.append(line)
    return "\n".join(lines)

def extract_text(file_path):
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":   return read_pdf(file_path)
    if ext == ".docx":  return read_docx(file_path)
    if ext == ".xlsx":  return read_xlsx(file_path)
    return read_txt(file_path)

# ─── PARSER ──────────────────────────────────────────────────────────────────

ANSWER_MAP = {"a": 0, "b": 1, "c": 2, "d": 3}

def parse_questions(text):
    questions = []
    blocks = re.split(r"\n(?=\s*\d+[\.\)]\s)", text)

    for block in blocks:
        block = block.strip()
        if not block:
            continue
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        if not lines:
            continue

        question_line = re.sub(r"^\d+[\.\)]\s*", "", lines[0]).strip()
        if not question_line:
            continue

        options = []
        correct_idx = None

        for line in lines[1:]:
            m = re.match(r"(?:to['\u2019]?g['\u2019]?ri|togri|javob|answer)\s*[:\-]\s*([a-dA-D])", line, re.IGNORECASE)
            if m:
                correct_idx = ANSWER_MAP.get(m.group(1).lower())
                continue
            m = re.match(r"^([a-dA-D])\s*[\.\)]\s*(.+)", line)
            if m:
                letter = m.group(1).lower()
                opt_text = m.group(2).strip()
                if re.search(r"[*+=]", opt_text):
                    correct_idx = ANSWER_MAP.get(letter)
                    opt_text = re.sub(r"\s*[*+=]\s*$", "", opt_text).strip()
                options.append(opt_text)

        if len(options) >= 2 and correct_idx is not None:
            correct_idx = min(correct_idx, len(options) - 1)
            questions.append({
                "question": question_line,
                "options": options,
                "correct_index": correct_idx
            })

    return questions

# ─── HELPERS ─────────────────────────────────────────────────────────────────

import uuid, json

def gen_id():
    return uuid.uuid4().hex[:8].upper()

def quiz_menu(quiz, bot_username="QuizBot"):
    count = len(quiz["questions"])
    link = f"https://t.me/{bot_username}?start=quiz_{quiz['quiz_id']}"
    kb = [
        [InlineKeyboardButton("▶️ Start quiz", callback_data=f"start:{quiz['quiz_id']}")],
        [InlineKeyboardButton("👥 Guruhda boshlash", callback_data=f"startgroup:{quiz['quiz_id']}")],
        [InlineKeyboardButton("📤 Ulashish havolasi", callback_data=f"share:{quiz['quiz_id']}")],
        [InlineKeyboardButton("✏️ Nom o'zgartirish", callback_data=f"edittitle:{quiz['quiz_id']}"),
         InlineKeyboardButton("🗑 O'chirish", callback_data=f"delete:{quiz['quiz_id']}")],
        [InlineKeyboardButton("📊 Statistika", callback_data=f"stats:{quiz['quiz_id']}")],
    ]
    text = (
        f"📋 *{quiz['title']}*\n"
        f"❓ {count} ta savol · ⏱ 30 soniya\n\n"
        f"🔗 Havola:\n`{link}`"
    )
    return text, InlineKeyboardMarkup(kb)

# ─── SEND NEXT QUESTION ───────────────────────────────────────────────────────

async def send_next_question(chat_id, session_id, app):
    session = get_session(session_id)
    if not session or not session[4]:  # not active
        return

    quiz = get_quiz(session[2])
    if not quiz:
        return

    q_index = session[3]
    questions = quiz["questions"]

    if q_index >= len(questions):
        close_session(session_id)
        rows = get_leaderboard(chat_id)
        medals = ["🥇", "🥈", "🥉"]
        lines = ["🎉 *Quiz tugadi!*\n\n🏆 *Natijalar:*\n"]
        for i, (fn, un, correct, total) in enumerate(rows[:5]):
            medal = medals[i] if i < 3 else f"{i+1}."
            pct = round(correct / total * 100) if total else 0
            name = f"@{un}" if un else fn
            lines.append(f"{medal} {name} — {correct}/{total} ({pct}%)")
        if not rows:
            lines.append("Hali hech kim javob bermadi.")
        await app.bot.send_message(chat_id, "\n".join(lines), parse_mode="Markdown")
        return

    q = questions[q_index]
    options = [str(o)[:100] for o in q["options"]]
    correct_idx = min(q["correct_index"], len(options) - 1)
    question_text = f"[{q_index+1}/{len(questions)}] {q['question']}"[:300]

    try:
        poll_msg = await app.bot.send_poll(
            chat_id=chat_id,
            question=question_text,
            options=options,
            type=Poll.QUIZ,
            correct_option_id=correct_idx,
            is_anonymous=False,
            open_period=30,
        )
        save_poll(poll_msg.poll.id, chat_id, correct_idx, session_id)
        update_session_index(session_id, q_index + 1)

        import asyncio
        async def delayed_next():
            await asyncio.sleep(35)
            await send_next_question(chat_id, session_id, app)
        asyncio.ensure_future(delayed_next())

    except Exception as e:
        logger.error(f"Poll xatosi: {e}")
        update_session_index(session_id, q_index + 1)
        import asyncio
        asyncio.ensure_future(send_next_question(chat_id, session_id, app))

# ─── HANDLERS ────────────────────────────────────────────────────────────────

FORMAT_MSG = (
    "📋 *Fayl formati:*\n\n"
    "```\n"
    "1. Savol matni\n"
    "a) Birinchi javob\n"
    "b) Ikkinchi javob\n"
    "c) To'g'ri javob\n"
    "d) To'rtinchi javob\n"
    "to'g'ri: c\n\n"
    "2. Keyingi savol...\n"
    "```\n\n"
    "✅ Yoki variant yoniga `*` qo'ying:\n"
    "`c) To'g'ri javob *`"
)

async def cmd_start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    args = ctx.args
    # Deep link orqali quiz boshlash: /start quiz_XXXXXXXX
    if args and args[0].startswith("quiz_"):
        quiz_id = args[0][5:]
        quiz = get_quiz(quiz_id)
        if not quiz:
            await update.message.reply_text("❌ Quiz topilmadi.")
            return
        text, kb = quiz_menu(quiz)
        await update.message.reply_text(text, reply_markup=kb, parse_mode="Markdown")
        return

    await update.message.reply_text(
        "👋 Salom! Men *QuizBot* man.\n\n"
        "📎 PDF, DOCX, TXT yoki XLSX fayl yuboring\n"
        "Men undan quiz yasab beraman!\n\n"
        "📌 /format — fayl formati\n"
        "📚 /quizzes — mening quizlarim\n"
        "🏆 /top — leaderboard\n"
        "⛔️ /stop — quizni to'xtatish",
        parse_mode="Markdown"
    )

async def cmd_format(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(FORMAT_MSG, parse_mode="Markdown")

async def cmd_quizzes(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    quizzes = get_user_quizzes(update.effective_user.id)
    if not quizzes:
        await update.message.reply_text("📚 Sizda hali quiz yo'q.\n\n📎 Fayl yuboring — men quiz yasayman!")
        return
    kb = []
    for q in quizzes:
        kb.append([InlineKeyboardButton(
            f"📋 {q['title']} ({q['count']} savol)",
            callback_data=f"view:{q['quiz_id']}"
        )])
    await update.message.reply_text(
        "📚 *Sizning quizlaringiz:*",
        reply_markup=InlineKeyboardMarkup(kb),
        parse_mode="Markdown"
    )

async def cmd_top(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    rows = get_leaderboard(update.effective_chat.id)
    if not rows:
        await update.message.reply_text("📊 Hali hech kim quiz yechmagan.")
        return
    medals = ["🥇", "🥈", "🥉"]
    lines = ["🏆 *Leaderboard*\n"]
    for i, (fn, un, correct, total) in enumerate(rows):
        medal = medals[i] if i < 3 else f"{i+1}."
        pct = round(correct / total * 100) if total else 0
        name = f"@{un}" if un else fn
        lines.append(f"{medal} {name} — {correct}/{total} ({pct}%)")
    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")

async def cmd_stop(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    close_all_sessions(chat_id)
    await update.message.reply_text("⛔️ Quiz to'xtatildi.")

async def cmd_reset(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    chat = update.effective_chat
    user = update.effective_user
    if chat.type in ("group", "supergroup"):
        admins = await ctx.bot.get_chat_administrators(chat.id)
        if user.id not in [a.user.id for a in admins]:
            await update.message.reply_text("❌ Faqat adminlar reset qila oladi.")
            return
    conn = sqlite3.connect("quiz.db")
    conn.execute("DELETE FROM scores WHERE chat_id=?", (chat.id,))
    conn.commit()
    conn.close()
    await update.message.reply_text("✅ Leaderboard tozalandi.")

async def handle_document(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    ext = Path(doc.file_name or "").suffix.lower()

    if ext not in (".pdf", ".docx", ".txt", ".xlsx"):
        await update.message.reply_text(
            "❌ Faqat PDF, DOCX, TXT yoki XLSX.\n📌 /format — fayl formati"
        )
        return

    msg = await update.message.reply_text("⏳ Fayl o'qilmoqda...")

    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp_path = tmp.name
    try:
        tg_file = await doc.get_file()
        await tg_file.download_to_drive(tmp_path)
        raw_text = extract_text(tmp_path)
    except Exception as e:
        await msg.edit_text(f"❌ Faylni o'qishda xato: {e}")
        return
    finally:
        Path(tmp_path).unlink(missing_ok=True)

    if not raw_text.strip():
        await msg.edit_text("❌ Fayl bo'sh yoki o'qib bo'lmadi.")
        return

    questions = parse_questions(raw_text)
    if not questions:
        await msg.edit_text(
            "❌ Savollar topilmadi.\n📌 /format — fayl qanday bo'lishi kerak"
        )
        return

    # Save quiz
    quiz_id = gen_id()
    title = Path(doc.file_name).stem[:40]
    save_quiz(quiz_id, update.effective_user.id, title, questions)

    quiz = get_quiz(quiz_id)
    bot_username = (await ctx.bot.get_me()).username
    text, kb = quiz_menu(quiz, bot_username)
    await msg.edit_text(
        f"✅ *{len(questions)} ta savol topildi!*\n\n{text}",
        reply_markup=kb,
        parse_mode="Markdown"
    )

async def handle_callback(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    data = q.data
    chat_id = q.message.chat_id

    if data.startswith("view:"):
        quiz_id = data[5:]
        quiz = get_quiz(quiz_id)
        if not quiz:
            await q.edit_message_text("❌ Quiz topilmadi.")
            return
        bot_username = (await ctx.bot.get_me()).username
        text, kb = quiz_menu(quiz, bot_username)
        await q.edit_message_text(text, reply_markup=kb, parse_mode="Markdown")

    elif data.startswith("start:"):
        quiz_id = data[6:]
        quiz = get_quiz(quiz_id)
        if not quiz:
            await q.edit_message_text("❌ Quiz topilmadi.")
            return
        close_all_sessions(chat_id)
        session_id = gen_id()
        create_session(session_id, chat_id, quiz_id)
        await q.edit_message_text(
            f"▶️ *{quiz['title']}* boshlanmoqda!\n"
            f"❓ {len(quiz['questions'])} ta savol • ⏱ 30 soniya",
            parse_mode="Markdown"
        )
        import asyncio
        asyncio.create_task(send_next_question(chat_id, session_id, ctx.application))

    elif data.startswith("startgroup:"):
        quiz_id = data[11:]
        bot_username = (await ctx.bot.get_me()).username
        link = f"https://t.me/{bot_username}?start=quiz_{quiz_id}"
        await q.answer(
            f"Guruhda boshlash uchun:\n{link}\nHavolani guruhga yuboring!",
            show_alert=True
        )

    elif data.startswith("share:"):
        quiz_id = data[6:]
        bot_username = (await ctx.bot.get_me()).username
        link = f"https://t.me/{bot_username}?start=quiz_{quiz_id}"
        await q.answer(f"🔗 Havola:\n{link}", show_alert=True)

    elif data.startswith("stats:"):
        quiz_id = data[6:]
        quiz = get_quiz(quiz_id)
        sessions = get_quiz_stats(quiz_id)
        await q.answer(
            f"📊 {quiz['title']}\n"
            f"❓ {len(quiz['questions'])} savol\n"
            f"▶️ {sessions} marta o'ynaldi",
            show_alert=True
        )

    elif data.startswith("delete:"):
        quiz_id = data[7:]
        kb = InlineKeyboardMarkup([[
            InlineKeyboardButton("✅ Ha, o'chir", callback_data=f"confirmdelete:{quiz_id}"),
            InlineKeyboardButton("❌ Yo'q", callback_data=f"view:{quiz_id}")
        ]])
        await q.edit_message_text("🗑 Rostdan ham o'chirmoqchimisiz?", reply_markup=kb)

    elif data.startswith("confirmdelete:"):
        quiz_id = data[14:]
        delete_quiz(quiz_id)
        await q.edit_message_text("✅ Quiz o'chirildi.")

    elif data.startswith("edittitle:"):
        quiz_id = data[10:]
        ctx.user_data["editing_title"] = quiz_id
        await q.edit_message_text(
            "✏️ Yangi nom yozing:",
            reply_markup=InlineKeyboardMarkup([[
                InlineKeyboardButton("❌ Bekor", callback_data=f"view:{quiz_id}")
            ]])
        )

async def handle_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    # Editing quiz title
    if "editing_title" in ctx.user_data:
        quiz_id = ctx.user_data.pop("editing_title")
        new_title = update.message.text.strip()[:40]
        conn = sqlite3.connect("quiz.db")
        conn.execute("UPDATE quizzes SET title=? WHERE quiz_id=?", (new_title, quiz_id))
        conn.commit()
        conn.close()
        quiz = get_quiz(quiz_id)
        text, kb = quiz_menu(quiz)
        await update.message.reply_text(
            f"✅ Nom o'zgartirildi!\n\n{text}",
            reply_markup=kb,
            parse_mode="Markdown"
        )

async def handle_poll_answer(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    answer = update.poll_answer
    row = get_poll(answer.poll_id)
    if not row or not answer.option_ids:
        return
    chat_id, correct_idx, session_id = row
    user = answer.user
    upsert_score(chat_id, user.id, user.username or "", user.first_name or "Nomsiz",
                 answer.option_ids[0] == correct_idx)

# ─── MAIN ────────────────────────────────────────────────────────────────────

def main():
    init_db()
    app = Application.builder().token(BOT_TOKEN).build()

    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("format", cmd_format))
    app.add_handler(CommandHandler("quizzes", cmd_quizzes))
    app.add_handler(CommandHandler("top", cmd_top))
    app.add_handler(CommandHandler("stop", cmd_stop))
    app.add_handler(CommandHandler("reset", cmd_reset))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_handler(PollAnswerHandler(handle_poll_answer))

    logger.info("Bot ishga tushdi ✅")
    app.run_polling(drop_pending_updates=True)

if __name__ == "__main__":
    main()
